from datetime import datetime, timedelta
from uuid import uuid4

import pytest

from app.models.reminder import Notification

# 测试库在整个会话内共享，公告与反馈都不按用户隔离，
# 因此每个用例用唯一标题并只校验自己创建的数据。
def tag() -> str:
    return uuid4().hex[:6]


def publish(client, admin_headers, **overrides):
    payload = {
        "title": f"国庆寄养须知-{tag()}",
        "category": "notice",
        "content": "国庆期间寄养需提前 7 天预约，请携带疫苗本。",
        "audience": "all",
        "status": "published",
        **overrides,
    }
    response = client.post("/api/admin/announcements", headers=admin_headers, json=payload)
    assert response.json()["code"] == 200, response.text
    return response.json()["data"]["id"], payload["title"]


def listing(client, headers):
    body = client.get("/api/announcements", headers=headers).json()
    assert body["code"] == 200, body
    return body["data"]["items"]


def find(items, announcement_id):
    return next((item for item in items if item["id"] == announcement_id), None)


def test_announcement_audience_unread_and_read_stats(
    client, admin_headers, owner_headers, staff_headers
):
    announcement_id, title = publish(client, admin_headers)
    assert find(listing(client, owner_headers), announcement_id)["is_read"] is False
    assert find(listing(client, staff_headers), announcement_id)["is_read"] is False

    detail = client.get(f"/api/announcements/{announcement_id}", headers=owner_headers).json()
    assert detail["code"] == 200
    assert "疫苗本" in detail["data"]["content"]
    assert find(listing(client, owner_headers), announcement_id)["is_read"] is True
    assert find(listing(client, staff_headers), announcement_id)["is_read"] is False

    # 重复打开不重复计入已读。
    client.get(f"/api/announcements/{announcement_id}", headers=owner_headers)
    stats = client.get(
        f"/api/admin/announcements/{announcement_id}/reads", headers=admin_headers
    ).json()["data"]
    assert stats["read"] == 1 == len(stats["readers"])
    assert stats["target"] >= 3
    assert stats["unread"] == stats["target"] - 1


def test_announcement_audience_targeting_and_pinning(
    client, admin_headers, owner_headers, staff_headers
):
    staff_id, staff_title = publish(client, admin_headers, audience="boarding_staff")
    pin_id, pin_title = publish(client, admin_headers, is_pinned=True)

    owner_items = listing(client, owner_headers)
    assert find(owner_items, staff_id) is None
    staff_items = listing(client, staff_headers)
    assert find(staff_items, staff_id) is not None

    # 置顶公告排在任何非置顶公告之前。
    ids = [item["id"] for item in owner_items]
    first_unpinned = next(item["id"] for item in owner_items if not item["is_pinned"])
    assert ids.index(pin_id) < ids.index(first_unpinned)

    client.post("/api/announcements/read-all", headers=owner_headers)
    assert all(item["is_read"] for item in listing(client, owner_headers))
    unread = client.get("/api/announcements/unread-count", headers=owner_headers).json()["data"]
    assert unread["unread"] == 0


def test_announcement_draft_and_expired_hidden(client, admin_headers, owner_headers):
    draft_id, _ = publish(client, admin_headers, status="draft")
    expired_id, _ = publish(
        client,
        admin_headers,
        publish_at=(datetime.now() - timedelta(days=5)).isoformat(),
        expire_at=(datetime.now() - timedelta(days=1)).isoformat(),
    )
    ids = [item["id"] for item in listing(client, owner_headers)]
    assert draft_id not in ids and expired_id not in ids
    assert (
        client.get(f"/api/announcements/{draft_id}", headers=owner_headers).json()["code"]
        == 404
    )
    assert (
        client.get(f"/api/announcements/{expired_id}", headers=owner_headers).json()["code"]
        == 404
    )


def test_announcement_admin_permissions_and_delete_rule(
    client, admin_headers, owner_headers, staff_headers
):
    payload = {"title": f"越权-{tag()}", "content": "内容", "status": "published"}
    assert client.post("/api/admin/announcements", headers=owner_headers, json=payload).json()["code"] == 403
    assert client.post("/api/admin/announcements", headers=staff_headers, json=payload).json()["code"] == 403

    announcement_id, title = publish(client, admin_headers)
    # 已发布不能直接删，先归档。
    assert client.delete(f"/api/admin/announcements/{announcement_id}", headers=admin_headers).json()["code"] == 409
    archived = client.put(
        f"/api/admin/announcements/{announcement_id}",
        headers=admin_headers,
        json={"title": title, "content": "已归档", "status": "archived"},
    )
    assert archived.json()["code"] == 200
    assert client.delete(f"/api/admin/announcements/{announcement_id}", headers=admin_headers).json()["code"] == 200

    # 上线时间晚于过期时间要被拒绝。
    invalid = client.post(
        "/api/admin/announcements",
        headers=admin_headers,
        json={
            "title": f"时间错乱-{tag()}",
            "content": "内容",
            "status": "published",
            "publish_at": (datetime.now() + timedelta(days=2)).isoformat(),
            "expire_at": (datetime.now() + timedelta(days=1)).isoformat(),
        },
    ).json()
    assert invalid["code"] == 400


def test_feedback_flow_notifies_author_and_admin(
    client, admin_headers, owner_headers, staff_headers, db
):
    marker = tag()
    created = client.post(
        "/api/feedback",
        headers=owner_headers,
        json={"category": "bug", "content": f"预约页日期选择器在手机上溢出 {marker}", "contact": "13800000000"},
    ).json()
    assert created["code"] == 200
    feedback_id = created["data"]["id"]
    assert created["data"]["status"] == "pending"

    mine = client.get("/api/feedback", headers=owner_headers).json()["data"]["items"]
    assert feedback_id in [item["id"] for item in mine]

    listed = client.get("/api/admin/feedbacks", headers=admin_headers, params={"status": "pending"}).json()["data"]
    assert feedback_id in [item["id"] for item in listed["items"]]
    assert next(item for item in listed["items"] if item["id"] == feedback_id)["username"]

    # 普通用户不能回复或改状态。
    assert client.put(
        f"/api/admin/feedbacks/{feedback_id}",
        headers=owner_headers,
        json={"status": "resolved"},
    ).json()["code"] == 403

    replied = client.put(
        f"/api/admin/feedbacks/{feedback_id}",
        headers=admin_headers,
        json={"status": "resolved", "reply": "已修复，请刷新页面。"},
    ).json()
    assert replied["code"] == 200 and replied["data"]["status"] == "resolved"

    notices = (
        db.query(Notification)
        .filter(Notification.title == "意见反馈已更新", Notification.link == "/help")
        .order_by(Notification.id.desc())
        .all()
    )
    latest = notices[0]
    assert "已修复" in latest.content

    # 同一状态重复提交不会重复通知（dedupe_key 幂等）。
    before = db.query(Notification).filter(Notification.title == "意见反馈已更新").count()
    client.put(
        f"/api/admin/feedbacks/{feedback_id}",
        headers=admin_headers,
        json={"status": "resolved"},
    )
    assert db.query(Notification).filter(Notification.title == "意见反馈已更新").count() == before

    # 处理过的反馈不能退回待处理。
    assert client.put(
        f"/api/admin/feedbacks/{feedback_id}",
        headers=admin_headers,
        json={"status": "pending"},
    ).json()["code"] == 409


def test_feedback_validation(client, owner_headers):
    response = client.post("/api/feedback", headers=owner_headers, json={"content": "太短"})
    assert response.status_code in (400, 422)


def test_statistics_export_admin_only(client, admin_headers, owner_headers):
    assert (
        client.get("/api/admin/statistics/export", headers=owner_headers).json()["code"]
        == 403
    )
    response = client.get("/api/admin/statistics/export?days=7", headers=admin_headers)
    assert response.status_code == 200
    assert response.content[:2] == b"PK"
    assert "spreadsheetml" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]

    invalid = client.get("/api/admin/statistics/export?days=15", headers=admin_headers).json()
    assert invalid["code"] == 400


@pytest.mark.parametrize("days", [7, 30])
def test_statistics_export_opens_with_expected_sheets(client, admin_headers, days):
    from io import BytesIO

    from openpyxl import load_workbook

    response = client.get(f"/api/admin/statistics/export?days={days}", headers=admin_headers)
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["汇总", "每日明细", "同期订单"]
    daily = book["每日明细"]
    assert daily.max_row == days + 1
    assert daily.cell(row=1, column=1).value == "日期"
